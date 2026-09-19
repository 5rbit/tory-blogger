import importlib.util, sys
from pathlib import Path
from openpyxl import load_workbook
from src.common import save_yaml

def _mod():
    spec = importlib.util.spec_from_file_location("export_excel", Path("scripts/export_excel.py")); m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m

def test_export_and_import_roundtrip(tmp_path, monkeypatch):
    m = _mod()
    monkeypatch.setattr(m, "data_dir", lambda: tmp_path)
    monkeypatch.setattr(m, "pipeline_config", lambda: {"topic": {"main": "등산"}, "schedule": {"posts_per_week": 3}})
    monkeypatch.setattr(m, "set_experiment", lambda e: None)
    save_yaml(tmp_path / "keywords" / "2026-W40.yaml", {"week": "2026-W40", "posts": [{"keyword": "북한산 단풍", "post_type": "course", "status": "pending", "angle": "", "volume": 5000, "competition": 0.6, "personal_ratio": 0.5}]})
    (tmp_path / "drafts").mkdir(); (tmp_path / "drafts" / "20261001-x.md").write_text("# 제목\n\n📌 **한줄요약**: a\n\n## 🥾 코스\n[사진: a]\n", encoding="utf-8")
    out = m.export("hiker")
    wb = load_workbook(out); ws = wb["📋 키워드계획"]
    assert ws.max_row == 2 and ws["C2"].value == "북한산 단풍" and wb["✍️ 초안"].max_row == 2
    ws["E2"] = "approved"; ws["F2"] = "우이동 코스로"; wb.save(out)
    assert m.import_status("hiker") == 2
    from src.common import load_yaml
    p = load_yaml(tmp_path / "keywords" / "2026-W40.yaml")["posts"][0]
    assert p["status"] == "approved" and p["angle"] == "우이동 코스로"
