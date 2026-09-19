#!/usr/bin/env python3
"""작업 데이터를 엑셀 한 파일로 정리한다 (data/<exp>/tracker.xlsx).
시트: 📋 키워드계획 · ✍️ 초안 · 🖼️ 코스자산 · 📊 주간리포트 · ⚙️ 설정
--import 로 '📋 키워드계획' 시트의 status 열(approved/rejected)을 YAML 로 되돌려 쓴다 → 엑셀에서 승인 가능.
사용: python scripts/export_excel.py --exp hiker            (내보내기)
      python scripts/export_excel.py --exp hiker --import   (엑셀의 승인 상태를 YAML 에 반영)
"""
import sys, re; from pathlib import Path; sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from datetime import datetime
import click
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from src.common import data_dir, get_logger, list_experiments, load_yaml, pipeline_config, save_yaml, set_experiment
log = get_logger("export_excel")

HEAD = PatternFill("solid", fgColor="1F3D2B"); HFONT = Font(bold=True, color="FFFFFF")
STATUS_FILL = {"approved": "DFF0D8", "rejected": "F2DEDE", "drafted": "FFF3CD", "published": "D9EDF7", "pending": "FFFFFF"}

def sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title); ws.append(header)
    for c in ws[1]:
        c.fill = HEAD; c.font = HFONT; c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths or [18] * len(header), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws

def keyword_rows(d: Path):
    rows = []
    for y in sorted((d / "keywords").glob("*.yaml")):
        plan = load_yaml(y)
        for i, p in enumerate(plan.get("posts", [])):
            ex = p.get("extra", {}) or {}
            rows.append([plan.get("week", y.stem), i + 1, p.get("keyword"), p.get("post_type"), p.get("status"), p.get("angle", ""),
                         p.get("volume", 0), round(float(p.get("competition", 0)), 2), round(float(p.get("personal_ratio", 0)), 2),
                         ex.get("pc", ""), ex.get("mobile", ""), ", ".join((ex.get("serp") or {}).get("top_titles", [])[:3]), y.name])
    return rows

def draft_rows(d: Path):
    rows = []
    for m in sorted((d / "drafts").glob("*.md")) if (d / "drafts").exists() else []:
        t = m.read_text(encoding="utf-8"); title = (t.splitlines() or [""])[0].lstrip("# ")
        body = re.sub(r"\s", "", re.sub(r"#\S+", "", t))
        rows.append([m.stem[:8], title, len(body), len(re.findall(r"^## ", t, flags=re.M)), len(re.findall(r"\[사진:|^!\[", t, flags=re.M)),
                     "✅" if "**한줄요약**" in t else "❌", "✅" if "🚌 대중교통" in t or "## 대중교통" in t else "❌",
                     "✅" if "이번 주말 현황" in t else "❌", len(re.findall(r"\[상품:", t)), m.name])
    return rows

def asset_rows(d: Path):
    rows = []
    for a in sorted((d / "assets").glob("*")) if (d / "assets").exists() else []:
        if a.is_dir():
            files = {p.name for p in a.iterdir()}
            rows.append([a.name] + ["✅" if f in files else "❌" for f in ["card.png", "profile.png", "radar.png", "timeline.md"]] +
                        [datetime.fromtimestamp(a.stat().st_mtime).strftime("%Y-%m-%d")])
    return rows

def report_rows(d: Path):
    rows = []
    for r in sorted((d / "reports").glob("*.md")) if (d / "reports").exists() else []:
        rows.append([r.stem, r.read_text(encoding="utf-8").count("\n"), r.name])
    return rows

def export(exp: str) -> Path:
    set_experiment(exp); d = data_dir(); cfg = pipeline_config()
    wb = Workbook(); wb.remove(wb.active)
    ws = sheet(wb, "📋 키워드계획", ["주차", "#", "키워드", "유형", "status", "관점", "월검색량", "경쟁도", "개인비율", "PC", "모바일", "상위노출 제목", "파일"],
               keyword_rows(d), [10, 4, 26, 10, 11, 40, 10, 8, 8, 8, 8, 60, 16])
    for row in ws.iter_rows(min_row=2):
        row[4].fill = PatternFill("solid", fgColor=STATUS_FILL.get(str(row[4].value), "FFFFFF"))
    sheet(wb, "✍️ 초안", ["날짜", "제목", "글자수", "소제목", "사진", "한줄요약", "대중교통", "현황블록", "상품슬롯", "파일"], draft_rows(d), [10, 40, 8, 8, 6, 8, 8, 8, 8, 40])
    sheet(wb, "🖼️ 코스자산", ["산/코스", "카드", "고도", "레이더", "타임라인", "생성일"], asset_rows(d), [24, 8, 8, 8, 10, 12])
    sheet(wb, "📊 주간리포트", ["주차", "줄수", "파일"], report_rows(d), [12, 8, 24])
    flat = [(k, str(v)) for k, v in cfg.items() if not isinstance(v, dict)] + [(f"{k}.{k2}", str(v2)) for k, v in cfg.items() if isinstance(v, dict) for k2, v2 in v.items()]
    sheet(wb, "⚙️ 설정", ["항목", "값"], flat, [32, 80])
    out = d / "tracker.xlsx"; out.parent.mkdir(parents=True, exist_ok=True); wb.save(out)
    log.info("엑셀 저장: %s (키워드 %d, 초안 %d)", out, ws.max_row - 1, len(draft_rows(d)))
    return out

def import_status(exp: str) -> int:
    set_experiment(exp); d = data_dir(); wb = load_workbook(d / "tracker.xlsx"); ws = wb["📋 키워드계획"]
    changed = 0; plans: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        week, idx, kw, _, status, angle, *_rest = row; fname = row[-1]
        if not fname:
            continue
        y = d / "keywords" / fname; plan = plans.setdefault(fname, load_yaml(y))
        posts = plan.get("posts", [])
        if 1 <= int(idx) <= len(posts):
            p = posts[int(idx) - 1]
            if status and status != p.get("status") and status in STATUS_FILL:
                p["status"] = status; changed += 1
            if angle and angle != p.get("angle"):
                p["angle"] = angle; changed += 1
    for fname, plan in plans.items():
        save_yaml(d / "keywords" / fname, plan)
    log.info("엑셀 → YAML 반영 %d건", changed); return changed

@click.command()
@click.option("--exp", default=None)
@click.option("--all", "run_all", is_flag=True)
@click.option("--import", "do_import", is_flag=True, help="엑셀의 status/관점 열을 YAML 에 반영")
def main(exp, run_all, do_import):
    for e in (list_experiments() if run_all else [exp]):
        (import_status if do_import else export)(e)

if __name__ == "__main__":
    main()
